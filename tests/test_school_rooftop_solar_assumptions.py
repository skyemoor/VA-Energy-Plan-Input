"""
test_school_rooftop_solar_assumptions.py

Run with: python3 -m pytest test_school_rooftop_solar_assumptions.py -v
"""
import re

import openpyxl
import pytest

import school_rooftop_solar_assumptions as s

VA_SLCOE_MODEL_XLSX_PATH = "/home/claude/work/lp_package/VA_SLCOE_Model.xlsx"


class TestCapacityAssumptionsAreExactlyWhatTheUserSet:
    """Direct user instruction, 2026-08-27: High=850kW, Middle=500kW, Elementary=250kW,
    4-hour battery duration. These are user decisions, not derived figures -- confirm the
    constants hold the exact values specified, not an approximation."""

    def test_high_school_capacity(self):
        assert s.HIGH_SCHOOL_SOLAR_CAPACITY_KW == 850

    def test_middle_school_capacity(self):
        assert s.MIDDLE_SCHOOL_SOLAR_CAPACITY_KW == 500

    def test_elementary_school_capacity(self):
        assert s.ELEMENTARY_SCHOOL_SOLAR_CAPACITY_KW == 250

    def test_battery_duration(self):
        assert s.BATTERY_DURATION_HOURS == 4.0


class TestRationaleSectionDoesNotDriveTheAssumptions:
    """The capacity assumptions are a user decision, not a formula output -- confirm the two
    sections are genuinely independent (changing rationale constants would not silently change
    the capacity assumptions, since Section 2 does not reference Section 1's own constants)."""

    def test_high_school_rate_not_equal_to_raw_direct_average(self):
        # The user's own 850 kW is close to but NOT identical to the ~844.8 kW direct-observation
        # average -- confirms 850 was a chosen round number, not a copied formula result.
        direct_avg = (
            s.RATIONALE_HUGUENOT_HS_RICHMOND_KW
            + s.RATIONALE_PATRICK_HENRY_HS_ROANOKE_KW
            + s.RATIONALE_WILLIAM_FLEMING_HS_ROANOKE_KW
        ) / 3
        assert s.HIGH_SCHOOL_SOLAR_CAPACITY_KW != direct_avg
        assert abs(s.HIGH_SCHOOL_SOLAR_CAPACITY_KW - direct_avg) < 10  # but genuinely close

    def test_rationale_summary_returns_real_recorded_data(self):
        summary = s.rationale_summary()
        assert summary["high_school_direct_observations_kw"]["Huguenot HS (Richmond)"] == 534.3
        assert summary["middle_school_direct_observation_kw"][
            "Locust Grove MS (Orange Co., flagged as possibly large)"
        ] == 945


class TestLocalitySolarKwComputation:
    def test_fairfax_matches_hand_calculation(self):
        # 142 elem * 250 + 23 mid * 500 + 29 high * 850 = 35,500 + 11,500 + 24,650 = 71,650
        result = s.locality_solar_kw_and_battery_kwh("Fairfax")
        assert result["elementary_kw"] == 35_500
        assert result["middle_kw"] == 11_500
        assert result["high_kw"] == 24_650
        assert result["total_kw"] == 71_650

    def test_chesapeake_matches_hand_calculation(self):
        # 28*250 + 10*500 + 7*850 = 7,000 + 5,000 + 5,950 = 17,950
        result = s.locality_solar_kw_and_battery_kwh("Chesapeake")
        assert result["total_kw"] == 17_950

    def test_unknown_locality_raises_rather_than_silently_returning_zero(self):
        # A typo in a locality name should fail loudly, not silently compute a zero total.
        with pytest.raises(KeyError):
            s.build_school_division("Fairfax County")  # not the exact key used ("Fairfax")

    def test_component_kw_sum_equals_total_kw_for_every_locality(self):
        for locality in s.SCHOOL_COUNTS_BY_LOCALITY:
            result = s.locality_solar_kw_and_battery_kwh(locality)
            assert (
                result["elementary_kw"] + result["middle_kw"] + result["high_kw"]
                == result["total_kw"]
            )


class TestSchoolTypeGroupAndSchoolDivisionClasses:
    """New OO-structure-specific behavior: locality + school-type granularity (not
    per-individual-school), per the user's own explicit strategic-level scoping decision."""

    def test_school_division_has_exactly_three_type_groups(self):
        division = s.build_school_division("Fairfax")
        assert len(division.type_groups) == 3

    def test_group_of_type_returns_correct_group(self):
        division = s.build_school_division("Fairfax")
        high_group = division.group_of_type(s.SchoolType.HIGH)
        assert high_group.count == 29
        assert high_group.school_type == s.SchoolType.HIGH

    def test_group_of_type_returns_none_for_absent_type(self):
        # No locality in this module has an absent type, but the method itself should degrade
        # gracefully rather than raise, since a future locality might genuinely have zero of a type.
        division = s.SchoolDivision(name="Test", type_groups=[])
        assert division.group_of_type(s.SchoolType.HIGH) is None

    def test_per_group_override_takes_precedence_over_flat_default(self):
        # Simulates a future, more granular pass setting one locality's own real figure.
        group = s.SchoolTypeGroup(
            locality="Richmond", school_type=s.SchoolType.HIGH, count=5,
            solar_kw_per_school_override=534.3,  # Huguenot's own real, sourced figure
        )
        assert group.solar_kw_per_school() == 534.3
        assert group.total_solar_kw() == 5 * 534.3  # override applies to the WHOLE group's count

    def test_override_does_not_affect_other_localities(self):
        # An override set on one SchoolTypeGroup instance must not leak into the flat module-level
        # default used by every other locality's own group.
        overridden = s.SchoolTypeGroup(
            "Richmond", s.SchoolType.HIGH, 5, solar_kw_per_school_override=534.3
        )
        untouched = s.build_school_division("Chesapeake").group_of_type(s.SchoolType.HIGH)
        assert overridden.solar_kw_per_school() == 534.3
        assert untouched.solar_kw_per_school() == s.HIGH_SCHOOL_SOLAR_CAPACITY_KW


class TestPlaceholderAreaFieldsForFutureWork:
    """The rooftop/parking-lot area fields exist as placeholders for the next planned step in
    this project's own work, but must not fabricate a kW estimate before a real density factor
    (kW/sqft) is established."""

    def test_rooftop_estimate_is_none_when_area_unset(self):
        group = s.SchoolTypeGroup("Arlington", s.SchoolType.HIGH, 9)
        assert group.rooftop_solar_kw_estimate() is None

    def test_parking_lot_estimate_is_none_when_area_unset(self):
        group = s.SchoolTypeGroup("Arlington", s.SchoolType.HIGH, 9)
        assert group.parking_lot_solar_kw_estimate() is None

    def test_rooftop_estimate_raises_not_fabricates_once_area_is_set(self):
        # Setting the area alone must NOT silently produce a number -- no density factor exists
        # yet, and inventing one would be a fabricated figure presented as if it were computed.
        group = s.SchoolTypeGroup(
            "Arlington", s.SchoolType.HIGH, 9, avg_rooftop_area_sqft_per_school=50_000
        )
        with pytest.raises(NotImplementedError):
            group.rooftop_solar_kw_estimate()

    def test_rooftop_and_parking_lot_are_independent_fields(self):
        # Confirms the two area types don't share state or a density factor by accident.
        group = s.SchoolTypeGroup(
            "Arlington", s.SchoolType.HIGH, 9, avg_rooftop_area_sqft_per_school=50_000
        )
        assert group.avg_parking_lot_area_sqft_per_school is None
        assert group.parking_lot_solar_kw_estimate() is None  # unaffected by rooftop being set


class TestBatteryKwhComputation:
    def test_battery_kwh_is_total_kw_times_duration(self):
        result = s.locality_solar_kw_and_battery_kwh("Richmond")
        assert result["battery_kwh"] == result["total_kw"] * 4.0

    def test_battery_duration_applied_uniformly_not_per_school_type(self):
        # Confirms the same 4-hour figure is used regardless of school-type mix -- a locality with
        # more high schools should NOT get a different battery-duration multiplier.
        arlington = s.locality_solar_kw_and_battery_kwh("Arlington")  # high-school-heavy mix
        chesapeake = s.locality_solar_kw_and_battery_kwh("Chesapeake")  # elem-heavy mix
        assert arlington["battery_kwh"] / arlington["total_kw"] == pytest.approx(4.0)
        assert chesapeake["battery_kwh"] / chesapeake["total_kw"] == pytest.approx(4.0)


class TestAllLocalitiesSummaryAndGrandTotal:
    def test_all_localities_summary_returns_eleven_entries(self):
        summaries = s.all_localities_summary()
        assert len(summaries) == 11

    def test_grand_total_school_counts_match_hand_sum(self):
        # Elementary counts: 26+11+142+57+62+19+25+18+33+55+28 = 476
        total = s.total_across_all_localities()
        assert total["total_elementary_schools"] == 476

    def test_grand_total_kw_equals_sum_of_locality_totals(self):
        total = s.total_across_all_localities()
        summed_manually = sum(
            s.locality_solar_kw_and_battery_kwh(loc)["total_kw"]
            for loc in s.SCHOOL_COUNTS_BY_LOCALITY
        )
        assert total["total_solar_kw"] == summed_manually

    def test_grand_total_kw_matches_hand_calculation(self):
        # 264,750 kW confirmed by hand: sum of all 11 localities' own totals
        total = s.total_across_all_localities()
        assert total["total_solar_kw"] == 264_750

    def test_grand_total_battery_kwh_matches_hand_calculation(self):
        # 264,750 kW * 4 hours = 1,059,000 kWh
        total = s.total_across_all_localities()
        assert total["total_battery_kwh"] == 1_059_000

    def test_grand_total_battery_kwh_equals_total_kw_times_duration(self):
        total = s.total_across_all_localities()
        assert total["total_battery_kwh"] == total["total_solar_kw"] * s.BATTERY_DURATION_HOURS

    def test_component_kw_totals_sum_to_grand_total_kw(self):
        total = s.total_across_all_localities()
        assert (
            total["total_elementary_kw"] + total["total_middle_kw"] + total["total_high_kw"]
            == total["total_solar_kw"]
        )

    def test_school_type_counts_sum_to_total_schools(self):
        total = s.total_across_all_localities()
        assert (
            total["total_elementary_schools"]
            + total["total_middle_schools"]
            + total["total_high_schools"]
            == total["total_schools_all_types"]
        )


class TestExcludedFacilitiesDocumentedNotSilentlyDropped:
    """Per this project's own established practice: edge-case facilities that don't map to one of
    the three flat rates are excluded from the tally, but must be documented, not silently lost."""

    def test_excluded_facilities_dict_covers_localities_with_known_edge_cases(self):
        excluded = s.excluded_facilities_by_locality()
        # These 6 localities had real, named edge-case facilities identified during research
        for locality in ["Alexandria", "Fairfax", "Loudoun", "Prince William", "Richmond", "Hampton"]:
            assert locality in excluded
            assert len(excluded[locality]) > 0

    def test_excluded_facilities_are_not_double_counted_in_main_tally(self):
        # Spot check: Fairfax's own SCHOOL_COUNTS_BY_LOCALITY entry (142, 23, 29) sums to 194,
        # well under the 264 total facilities found in research (which includes the 70 excluded
        # admin/alt/special-ed/secondary facilities) -- confirms exclusions were genuinely applied,
        # not just documented in text while silently included in the numbers.
        elem, mid, high = s.SCHOOL_COUNTS_BY_LOCALITY["Fairfax"]
        assert elem + mid + high == 194
        assert elem + mid + high < 264  # the full, unfiltered facility count found in research


class TestXlsxAssumptionsTabStaysInSyncWithCode:
    """This project's own standing convention: every assumption is connected bidirectionally
    with code -- the XLSX 'Assumptions & Sources' tab is a manual mirror of the constants below,
    not a live sync. These tests are the enforcement mechanism: if a constant changes here without
    the matching XLSX row being updated by hand, these tests catch the drift rather than letting
    it accumulate silently. See VA_SLCOE_Model.xlsx rows 120-124 and this module's own inline
    "XLSX mirror: row N" comments next to each constant."""

    @staticmethod
    @pytest.fixture(scope="class")
    def assumptions_sheet():
        wb = openpyxl.load_workbook(VA_SLCOE_MODEL_XLSX_PATH, data_only=False)
        return wb["Assumptions & Sources"]

    def test_row_120_matches_high_school_capacity_constant(self, assumptions_sheet):
        value_cell = assumptions_sheet.cell(120, 3).value  # e.g. "850 kW"
        xlsx_kw = float(re.search(r"[\d.]+", value_cell).group())
        assert xlsx_kw == s.HIGH_SCHOOL_SOLAR_CAPACITY_KW

    def test_row_121_matches_middle_school_capacity_constant(self, assumptions_sheet):
        value_cell = assumptions_sheet.cell(121, 3).value
        xlsx_kw = float(re.search(r"[\d.]+", value_cell).group())
        assert xlsx_kw == s.MIDDLE_SCHOOL_SOLAR_CAPACITY_KW

    def test_row_122_matches_elementary_school_capacity_constant(self, assumptions_sheet):
        value_cell = assumptions_sheet.cell(122, 3).value
        xlsx_kw = float(re.search(r"[\d.]+", value_cell).group())
        assert xlsx_kw == s.ELEMENTARY_SCHOOL_SOLAR_CAPACITY_KW

    def test_row_123_matches_battery_duration_constant(self, assumptions_sheet):
        value_cell = assumptions_sheet.cell(123, 3).value  # e.g. "4.0 hours"
        xlsx_hours = float(re.search(r"[\d.]+", value_cell).group())
        assert xlsx_hours == s.BATTERY_DURATION_HOURS

    def test_row_124_matches_computed_grand_total(self, assumptions_sheet):
        # Row 124's value cell is a free-text summary ("264,750 kW solar / 1,059,000 kWh
        # battery") -- extract both numbers and compare against the live computed total, so if
        # anything upstream changes (a locality's own school count, a capacity constant), this
        # test fails rather than leaving a stale total sitting in the spreadsheet.
        value_cell = assumptions_sheet.cell(124, 3).value
        numbers = [int(n.replace(",", "")) for n in re.findall(r"[\d,]+", value_cell)]
        xlsx_kw, xlsx_kwh = numbers[0], numbers[1]
        total = s.total_across_all_localities()
        assert xlsx_kw == total["total_solar_kw"]
        assert xlsx_kwh == total["total_battery_kwh"]

    def test_section_header_present_on_all_five_new_rows(self, assumptions_sheet):
        expected_header = (
            "School Rooftop Solar -- per-type capacity & battery duration assumptions "
            "(rough estimates)"
        )
        for row in range(120, 125):
            assert assumptions_sheet.cell(row, 1).value == expected_header


class TestSchoolCountsByLocalityStructure:
    def test_exactly_eleven_localities_present(self):
        assert len(s.SCHOOL_COUNTS_BY_LOCALITY) == 11

    def test_every_locality_has_three_element_tuple(self):
        for locality, counts in s.SCHOOL_COUNTS_BY_LOCALITY.items():
            assert len(counts) == 3, f"{locality} should have (elementary, middle, high)"

    def test_prince_william_uses_corrected_nces_figure_not_stale_raw_notes_figure(self):
        # This session's own raw research notes file has an OLDER, superseded PWCS figure
        # (62 elem, 18 mid, 13-or-16 high, flagged as unresolved). The corrected figure -- from a
        # full, direct NCES federal database count -- is 62/17/13. Confirms the corrected figure
        # was used, not the stale one still sitting in the raw notes file.
        elem, mid, high = s.SCHOOL_COUNTS_BY_LOCALITY["Prince William"]
        assert (elem, mid, high) == (62, 17, 13)
